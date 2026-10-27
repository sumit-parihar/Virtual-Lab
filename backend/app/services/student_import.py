"""Student file import for the admin console: reads a CSV or Excel file and checks every row.

Expected columns (header names are not case sensitive): URN, Name, Email.
"""
import csv
import io
import uuid
from dataclasses import dataclass
from typing import Optional

from email_validator import EmailNotValidError, validate_email
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.models import Batch, Enrollment, User, UserRole

_COLUMN_NAMES = {
    "urn": {"urn", "roll no", "roll no.", "roll number", "roll_number", "rollno"},
    "name": {"name", "student name", "full name", "full_name"},
    "email": {"email", "email address", "e-mail"},
}


@dataclass
class ParsedRow:
    row: int
    urn: str
    name: str
    email: str


@dataclass
class Evaluated:
    row: ParsedRow
    action: str  # create | enroll | skip | invalid
    reason: Optional[str] = None


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _table_from_upload(filename: str, content: bytes) -> list[list[str]]:
    name = (filename or "").lower()
    if name.endswith(".xlsx"):
        try:
            sheet = load_workbook(io.BytesIO(content), read_only=True, data_only=True).active
            return [[_cell(c) for c in row] for row in sheet.iter_rows(values_only=True)]
        except Exception as exc:
            raise ValueError("The Excel file could not be read. Please upload a valid .xlsx file.") from exc
    if name.endswith(".csv"):
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = content.decode("latin-1")
        return [[_cell(c) for c in row] for row in csv.reader(io.StringIO(text))]
    raise ValueError("Unsupported file type. Please upload a .csv or .xlsx file.")


def parse_upload(filename: str, content: bytes, max_rows: int) -> list[ParsedRow]:
    """Return the data rows. Raises ValueError with a message for the admin if the file cannot be used."""
    table = _table_from_upload(filename, content)
    if not table:
        raise ValueError("The file is empty.")

    header = [h.lower() for h in table[0]]
    index = {}
    for field, names in _COLUMN_NAMES.items():
        found = [i for i, h in enumerate(header) if h in names]
        if not found:
            raise ValueError("The file must have the columns URN, Name and Email in its first row.")
        index[field] = found[0]

    rows = []
    for number, cells in enumerate(table[1:], start=2):
        if not any(cells):
            continue  # blank line
        get = lambda f: cells[index[f]] if index[f] < len(cells) else ""  # noqa: E731
        rows.append(ParsedRow(number, get("urn"), get("name"), get("email")))
    if not rows:
        raise ValueError("The file has no student rows.")
    if len(rows) > max_rows:
        raise ValueError(f"Too many rows ({len(rows)}). Please upload at most {max_rows} students at a time.")
    return rows


def evaluate_rows(db: Session, batch: Batch, rows: list[ParsedRow]) -> list[Evaluated]:
    """Decide what would happen to every row, without changing the database."""
    emails = {r.email.lower() for r in rows if r.email}
    users = {u.email: u for u in db.query(User).filter(User.email.in_(emails)).all()} if emails else {}
    urns = {r.urn for r in rows if r.urn}
    urn_owner = {u.roll_number: u for u in db.query(User).filter(User.roll_number.in_(urns)).all()} if urns else {}

    enrollment_batch = {}
    if users:
        pairs = (
            db.query(Enrollment.student_id, Batch.id, Batch.name, Batch.academic_year)
            .join(Batch, Batch.id == Enrollment.batch_id)
            .filter(Enrollment.student_id.in_([u.id for u in users.values()]))
            .all()
        )
        enrollment_batch = {sid: (bid, f"{name} ({year})") for sid, bid, name, year in pairs}

    seen_email: dict[str, int] = {}
    seen_urn: dict[str, int] = {}
    result: list[Evaluated] = []

    for r in rows:
        if not r.urn or not r.name or not r.email:
            missing = [label for label, value in (("URN", r.urn), ("Name", r.name), ("Email", r.email)) if not value]
            result.append(Evaluated(r, "invalid", "Missing " + ", ".join(missing)))
            continue
        try:
            email = validate_email(r.email, check_deliverability=False).normalized.lower()
        except EmailNotValidError:
            result.append(Evaluated(r, "invalid", "Invalid email address"))
            continue
        r.email = email

        if email in seen_email:
            result.append(Evaluated(r, "invalid", f"Duplicate email (also in row {seen_email[email]})"))
            continue
        if r.urn in seen_urn:
            result.append(Evaluated(r, "invalid", f"Duplicate URN (also in row {seen_urn[r.urn]})"))
            continue
        seen_email[email] = r.row
        seen_urn[r.urn] = r.row

        existing = users.get(email)
        if existing is not None:
            if existing.role != UserRole.student:
                result.append(Evaluated(r, "invalid", "This email belongs to an instructor or admin account"))
            elif existing.id in enrollment_batch:
                batch_id, label = enrollment_batch[existing.id]
                if batch_id == batch.id:
                    result.append(Evaluated(r, "skip", "Already in this batch"))
                else:
                    result.append(Evaluated(r, "skip", f"Already in batch {label}. Use Move student."))
            elif not existing.is_active:
                result.append(Evaluated(r, "invalid", "This account is deactivated"))
            else:
                result.append(Evaluated(r, "enroll", "Existing account will be added to this batch"))
            continue

        if r.urn in urn_owner:
            result.append(Evaluated(r, "invalid", "URN already used by another student"))
            continue
        result.append(Evaluated(r, "create"))
    return result


def template_workbook() -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Students"
    ws.append(["URN", "Name", "Email"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 36
    ws.column_dimensions["C"].width = 36
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()

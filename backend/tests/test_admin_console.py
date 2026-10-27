"""Admin console API: admin only, batches, instructors, students, import, account actions."""
import csv
import io
import uuid

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy.exc import IntegrityError

from app.models import Batch, Enrollment, InstructorAssignment, User, UserRole
from helpers import make_experiment, make_user


@pytest.fixture
def sent(monkeypatch):
    """Capture the invitation emails instead of sending them (the real sender uses the real database)."""
    outbox = []
    monkeypatch.setattr("app.routers.admin.send_credentials_emails", lambda recipients: outbox.extend(recipients))
    return outbox


def _batch(db, name=None, year="2026-27"):
    batch = Batch(name=name or f"B-{uuid.uuid4().hex[:6]}", academic_year=year)
    db.add(batch)
    db.commit()
    return batch.id


def _student(db, batch_id=None, name="Asha Patil", urn=None, email=None):
    user = make_user(db, UserRole.student)
    row = db.get(User, user.id)
    row.full_name, row.roll_number = name, urn or f"U{uuid.uuid4().hex[:6]}"
    if email:
        row.email = email
    if batch_id:
        db.add(Enrollment(student_id=user.id, batch_id=batch_id))
    db.commit()
    return user


def _csv(rows, header=("URN", "Name", "Email")):
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(header)
    writer.writerows(rows)
    return buffer.getvalue().encode()


def _upload(client, admin, batch_id, kind, content, filename="students.csv"):
    return client.post(
        f"/admin/batches/{batch_id}/students/{kind}",
        files={"file": (filename, content, "application/octet-stream")},
        headers=admin.headers,
    )


# ── Access ────────────────────────────────────────────────────────────────────

def test_admin_endpoints_are_admin_only(client, db):
    admin = make_user(db, UserRole.admin)
    instructor = make_user(db, UserRole.instructor)
    student = make_user(db, UserRole.student)
    batch_id = _batch(db)
    paths = ["/admin/batches", f"/admin/batches/{batch_id}", "/admin/students", "/admin/instructors", "/admin/students/template"]
    for path in paths:
        assert client.get(path, headers=admin.headers).status_code == 200, path
        assert client.get(path, headers=instructor.headers).status_code == 403, path
        assert client.get(path, headers=student.headers).status_code == 403, path
        assert client.get(path).status_code == 401, path


def test_old_batch_and_instructor_endpoints_are_gone(client, db):
    admin = make_user(db, UserRole.admin)
    assert client.get("/batches", headers=admin.headers).status_code == 404
    assert client.post("/auth/instructors", json={}, headers=admin.headers).status_code in (404, 405)


# ── Batches ───────────────────────────────────────────────────────────────────

def test_create_batch_without_instructor_then_list_it(client, db):
    admin = make_user(db, UserRole.admin)
    resp = client.post("/admin/batches", json={"name": "TE-CSE-A", "academic_year": "2026-27"}, headers=admin.headers)
    assert resp.status_code == 201
    body = resp.json()
    assert body["instructor"] is None and body["student_count"] == 0 and body["is_archived"] is False
    dup = client.post("/admin/batches", json={"name": "TE-CSE-A", "academic_year": "2026-27"}, headers=admin.headers)
    assert dup.status_code == 409


def test_batch_list_shows_instructor_and_student_count(client, db):
    admin = make_user(db, UserRole.admin)
    instructor = make_user(db, UserRole.instructor)
    batch_id = _batch(db)
    db.add(InstructorAssignment(instructor_id=instructor.id, batch_id=batch_id))
    _student(db, batch_id)
    _student(db, batch_id)
    row = next(b for b in client.get("/admin/batches", headers=admin.headers).json() if b["id"] == str(batch_id))
    assert row["student_count"] == 2 and row["instructor"]["id"] == str(instructor.id)


def test_replace_and_remove_the_batch_instructor(client, db):
    admin = make_user(db, UserRole.admin)
    old, new = make_user(db, UserRole.instructor), make_user(db, UserRole.instructor)
    teaching_other = _batch(db)
    db.add(InstructorAssignment(instructor_id=new.id, batch_id=teaching_other))  # new one already has a batch
    batch_id = _batch(db)
    url = f"/admin/batches/{batch_id}/instructor"

    assert client.put(url, json={"instructor_id": str(old.id)}, headers=admin.headers).json()["instructor"]["id"] == str(old.id)
    assert client.put(url, json={"instructor_id": str(new.id)}, headers=admin.headers).json()["instructor"]["id"] == str(new.id)
    assert len(client.get("/instructor/batches", headers=new.headers).json()) == 2  # keeps the other batch
    assert client.get("/instructor/batches", headers=old.headers).json() == []

    assert client.delete(url, headers=admin.headers).json()["instructor"] is None
    assert db.query(InstructorAssignment).filter_by(batch_id=batch_id).count() == 0


def test_cannot_assign_a_deactivated_instructor_or_a_non_instructor(client, db):
    admin = make_user(db, UserRole.admin)
    instructor = make_user(db, UserRole.instructor)
    student = make_user(db, UserRole.student)
    batch_id = _batch(db)
    url = f"/admin/batches/{batch_id}/instructor"
    assert client.put(url, json={"instructor_id": str(student.id)}, headers=admin.headers).status_code == 404
    client.post(f"/admin/users/{instructor.id}/deactivate", headers=admin.headers)
    assert client.put(url, json={"instructor_id": str(instructor.id)}, headers=admin.headers).status_code == 409


def test_archived_batch_is_hidden_and_can_be_restored(client, db):
    admin = make_user(db, UserRole.admin)
    instructor = make_user(db, UserRole.instructor)
    batch_id = _batch(db)
    db.add(InstructorAssignment(instructor_id=instructor.id, batch_id=batch_id))
    db.commit()

    assert client.post(f"/admin/batches/{batch_id}/archive", headers=admin.headers).json()["is_archived"] is True
    assert str(batch_id) not in {b["id"] for b in client.get("/admin/batches", headers=admin.headers).json()}
    assert str(batch_id) in {b["id"] for b in client.get("/admin/batches?include_archived=true", headers=admin.headers).json()}
    # Hidden from the instructor console too.
    assert client.get("/instructor/batches", headers=instructor.headers).json() == []
    assert client.get(f"/instructor/batches/{batch_id}", headers=instructor.headers).status_code == 404

    client.post(f"/admin/batches/{batch_id}/restore", headers=admin.headers)
    assert len(client.get("/instructor/batches", headers=instructor.headers).json()) == 1


def test_edit_batch_name_and_year(client, db):
    admin = make_user(db, UserRole.admin)
    batch_id = _batch(db, "Old", "2025-26")
    taken = _batch(db, "Taken", "2026-27")
    ok = client.patch(f"/admin/batches/{batch_id}", json={"name": "New", "academic_year": "2026-27"}, headers=admin.headers)
    assert ok.status_code == 200 and ok.json()["name"] == "New"
    conflict = client.patch(f"/admin/batches/{batch_id}", json={"name": "Taken"}, headers=admin.headers)
    assert conflict.status_code == 409 and taken


# ── Instructors and invitations ───────────────────────────────────────────────

def test_create_instructor_gets_a_temporary_password_and_must_change_it(client, db, sent):
    admin = make_user(db, UserRole.admin)
    resp = client.post("/admin/instructors", json={"full_name": "Dr Rao", "email": "Rao@College.edu"}, headers=admin.headers)
    assert resp.status_code == 201
    body = resp.json()
    assert body["email"] == "rao@college.edu" and body["account_state"] == "invited" and body["role"] == "instructor"

    assert len(sent) == 1 and sent[0]["email"] == "rao@college.edu"
    login = client.post("/auth/login", json={"email": "rao@college.edu", "password": sent[0]["password"]})
    assert login.status_code == 200 and login.json()["must_change_password"] is True

    again = client.post("/admin/instructors", json={"full_name": "Other", "email": "rao@college.edu"}, headers=admin.headers)
    assert again.status_code == 409 and len(sent) == 1


def test_edit_instructor_details_and_conflicts(client, db, sent):
    admin = make_user(db, UserRole.admin)
    one = client.post("/admin/instructors", json={"full_name": "One", "email": "one@college.edu"}, headers=admin.headers).json()
    client.post("/admin/instructors", json={"full_name": "Two", "email": "two@college.edu"}, headers=admin.headers)
    ok = client.patch(f"/admin/instructors/{one['id']}", json={"full_name": "Uno", "email": "uno@college.edu"}, headers=admin.headers)
    assert ok.status_code == 200 and ok.json()["full_name"] == "Uno"
    bad = client.patch(f"/admin/instructors/{one['id']}", json={"email": "two@college.edu"}, headers=admin.headers)
    assert bad.status_code == 409


def test_resend_invitation_issues_a_new_password_and_creates_no_account(client, db, sent):
    admin = make_user(db, UserRole.admin)
    created = client.post("/admin/instructors", json={"full_name": "Dr Rao", "email": "rao@college.edu"}, headers=admin.headers).json()
    first = sent[0]["password"]
    users_before = db.query(User).count()

    resp = client.post(f"/admin/users/{created['id']}/resend-invitation", headers=admin.headers)
    assert resp.status_code == 200 and resp.json()["account_state"] == "invited"
    assert db.query(User).count() == users_before
    assert len(sent) == 2 and sent[1]["password"] != first
    assert client.post("/auth/login", json={"email": "rao@college.edu", "password": first}).status_code == 401
    assert client.post("/auth/login", json={"email": "rao@college.edu", "password": sent[1]["password"]}).status_code == 200

    # Admin accounts cannot be invited; deactivated ones must be reactivated first.
    assert client.post(f"/admin/users/{admin.id}/resend-invitation", headers=admin.headers).status_code == 404
    client.post(f"/admin/users/{created['id']}/deactivate", headers=admin.headers)
    assert client.post(f"/admin/users/{created['id']}/resend-invitation", headers=admin.headers).status_code == 409


def test_account_state_and_invitation_status(client, db):
    admin = make_user(db, UserRole.admin)
    student = _student(db)
    row = db.get(User, student.id)
    row.must_change_password, row.credentials_status = True, "failed"
    db.commit()
    shown = next(s for s in client.get("/admin/students", headers=admin.headers).json() if s["id"] == str(student.id))
    assert shown["account_state"] == "invited" and shown["invitation"] == "failed"

    row.must_change_password = False
    db.commit()
    shown = next(s for s in client.get("/admin/students", headers=admin.headers).json() if s["id"] == str(student.id))
    assert shown["account_state"] == "active"


# ── Deactivate / reactivate ───────────────────────────────────────────────────

def test_deactivated_user_cannot_log_in_and_can_be_reactivated(client, db):
    admin = make_user(db, UserRole.admin)
    student = _student(db, email="asha@college.edu")
    assert client.post("/auth/login", json={"email": "asha@college.edu", "password": "Pass12345"}).status_code == 200

    resp = client.post(f"/admin/users/{student.id}/deactivate", headers=admin.headers)
    assert resp.json()["account_state"] == "deactivated"
    assert client.post("/auth/login", json={"email": "asha@college.edu", "password": "Pass12345"}).status_code == 401

    client.post(f"/admin/users/{student.id}/reactivate", headers=admin.headers)
    assert client.post("/auth/login", json={"email": "asha@college.edu", "password": "Pass12345"}).status_code == 200


def test_admin_cannot_deactivate_themselves(client, db):
    admin = make_user(db, UserRole.admin)
    assert client.post(f"/admin/users/{admin.id}/deactivate", headers=admin.headers).status_code == 409
    assert db.get(User, admin.id).is_active is True


def test_deactivating_an_instructor_keeps_the_batch_assigned(client, db):
    admin = make_user(db, UserRole.admin)
    instructor = make_user(db, UserRole.instructor)
    batch_id = _batch(db)
    db.add(InstructorAssignment(instructor_id=instructor.id, batch_id=batch_id))
    db.commit()
    resp = client.post(f"/admin/users/{instructor.id}/deactivate", headers=admin.headers).json()
    assert [b["id"] for b in resp["batches"]] == [str(batch_id)]  # the admin can see it still needs a new instructor
    assert db.query(InstructorAssignment).filter_by(batch_id=batch_id).count() == 1


# ── Students: edit and move ───────────────────────────────────────────────────

def test_edit_student_details_and_reject_conflicts(client, db):
    admin = make_user(db, UserRole.admin)
    one = _student(db, urn="U1", email="one@college.edu")
    _student(db, urn="U2", email="two@college.edu")
    ok = client.patch(f"/admin/students/{one.id}", json={"full_name": "New Name", "roll_number": "U9", "email": "new@college.edu"}, headers=admin.headers)
    assert ok.status_code == 200
    assert (ok.json()["full_name"], ok.json()["roll_number"], ok.json()["email"]) == ("New Name", "U9", "new@college.edu")
    assert client.patch(f"/admin/students/{one.id}", json={"roll_number": "U2"}, headers=admin.headers).status_code == 409
    assert client.patch(f"/admin/students/{one.id}", json={"email": "two@college.edu"}, headers=admin.headers).status_code == 409


def test_move_student_carries_progress_to_the_new_batch(client, db):
    admin = make_user(db, UserRole.admin)
    old_teacher, new_teacher = make_user(db, UserRole.instructor), make_user(db, UserRole.instructor)
    old_batch, new_batch = _batch(db), _batch(db)
    db.add_all([
        InstructorAssignment(instructor_id=old_teacher.id, batch_id=old_batch),
        InstructorAssignment(instructor_id=new_teacher.id, batch_id=new_batch),
    ])
    student = _student(db, old_batch)
    exp = make_experiment(db)
    assert client.post(f"/experiments/{exp.slug}/sections/overview/complete", headers=student.headers).status_code == 200

    moved = client.post(f"/admin/students/{student.id}/move", json={"batch_id": str(new_batch)}, headers=admin.headers)
    assert moved.status_code == 200 and moved.json()["batch"]["id"] == str(new_batch)
    assert db.query(Enrollment).filter_by(student_id=student.id).count() == 1

    old_ids = {s["id"] for s in client.get(f"/instructor/batches/{old_batch}", headers=old_teacher.headers).json()["students"]}
    assert str(student.id) not in old_ids
    rows = client.get(f"/instructor/batches/{new_batch}/experiments/{exp.id}/status", headers=new_teacher.headers).json()["rows"]
    assert [r["status"] for r in rows] == ["in_progress"]  # their progress moved with them


def test_cannot_move_a_student_into_an_archived_batch(client, db):
    admin = make_user(db, UserRole.admin)
    student = _student(db, _batch(db))
    archived = _batch(db)
    client.post(f"/admin/batches/{archived}/archive", headers=admin.headers)
    assert client.post(f"/admin/students/{student.id}/move", json={"batch_id": str(archived)}, headers=admin.headers).status_code == 409


def test_the_database_allows_only_one_batch_per_student(db):
    student = _student(db, _batch(db))
    db.add(Enrollment(student_id=student.id, batch_id=_batch(db)))
    with pytest.raises(IntegrityError):
        db.commit()
    db.rollback()


# ── Student import ────────────────────────────────────────────────────────────

def test_template_is_an_excel_file_with_the_three_columns(client, db):
    admin = make_user(db, UserRole.admin)
    resp = client.get("/admin/students/template", headers=admin.headers)
    assert resp.status_code == 200 and "spreadsheetml" in resp.headers["content-type"]
    ws = load_workbook(io.BytesIO(resp.content)).active
    assert [c.value for c in ws[1]] == ["URN", "Name", "Email"]


def test_preview_checks_every_row_and_changes_nothing(client, db):
    admin = make_user(db, UserRole.admin)
    batch_id, other_batch = _batch(db, "A"), _batch(db, "B")
    _student(db, other_batch, urn="X1", email="elsewhere@college.edu")
    _student(db, batch_id, urn="X2", email="here@college.edu")
    _student(db, None, urn="TAKEN", email="someone@college.edu")
    teacher = make_user(db, UserRole.instructor)
    db.get(User, teacher.id).email = "teacher@college.edu"
    db.commit()
    users_before = db.query(User).count()

    content = _csv([
        ("U1", "Good One", "good1@college.edu"),
        ("U2", "Good Two", "good2@college.edu"),
        ("U3", "Bad Email", "not-an-email"),
        ("U4", "", "noname@college.edu"),
        ("U5", "Dup Email", "good1@college.edu"),
        ("U1", "Dup Urn", "dupurn@college.edu"),
        ("U6", "Elsewhere", "elsewhere@college.edu"),
        ("U7", "Here", "here@college.edu"),
        ("TAKEN", "Urn Taken", "newperson@college.edu"),
        ("U8", "A Teacher", "teacher@college.edu"),
    ])
    body = _upload(client, admin, batch_id, "preview", content).json()

    by_row = {r["row"]: r for r in body["rows"]}
    assert [by_row[2]["status"], by_row[3]["status"]] == ["ready", "ready"]
    assert "Invalid email" in by_row[4]["reason"]
    assert "Missing Name" in by_row[5]["reason"]
    assert "Duplicate email" in by_row[6]["reason"]
    assert "Duplicate URN" in by_row[7]["reason"]
    assert by_row[8]["status"] == "skipped" and "Use Move student" in by_row[8]["reason"]
    assert by_row[9]["status"] == "skipped" and "Already in this batch" in by_row[9]["reason"]
    assert "URN already used" in by_row[10]["reason"]
    assert "instructor or admin" in by_row[11]["reason"]
    assert (body["ready"], body["skipped"], body["invalid"], body["total"]) == (2, 2, 6, 10)
    assert body["preview"] is True
    assert db.query(User).count() == users_before  # nothing was created


def test_import_creates_the_good_rows_and_emails_only_new_accounts(client, db, sent):
    admin = make_user(db, UserRole.admin)
    batch_id = _batch(db)
    content = _csv([("U1", "Asha Patil", "Asha@College.edu"), ("U2", "Bad", "oops"), ("U3", "Ravi Kumar", "ravi@college.edu")])

    body = _upload(client, admin, batch_id, "import", content).json()
    assert (body["created"], body["invalid"], body["skipped"]) == (2, 1, 0)
    assert [r["status"] for r in body["rows"]] == ["created", "invalid", "created"]

    asha = db.query(User).filter_by(email="asha@college.edu").one()
    assert asha.role == UserRole.student and asha.roll_number == "U1" and asha.must_change_password is True
    assert db.query(Enrollment).filter_by(batch_id=batch_id).count() == 2
    assert sorted(r["email"] for r in sent) == ["asha@college.edu", "ravi@college.edu"]
    assert client.post("/auth/login", json={"email": "asha@college.edu", "password": sent[0]["password"]}).status_code == 200

    # Importing the same file again creates nothing new and sends no more emails.
    again = _upload(client, admin, batch_id, "import", content).json()
    assert again["created"] == 0 and again["skipped"] == 2
    assert len(sent) == 2


def test_import_reads_excel_files_too(client, db, sent):
    admin = make_user(db, UserRole.admin)
    batch_id = _batch(db)
    wb = Workbook()
    ws = wb.active
    ws.append(["urn", "NAME", "Email"])  # header case does not matter
    ws.append([1001, "Excel Student", "excel@college.edu"])  # numeric URN cell
    buffer = io.BytesIO()
    wb.save(buffer)
    body = _upload(client, admin, batch_id, "import", buffer.getvalue(), "students.xlsx").json()
    assert body["created"] == 1
    assert db.query(User).filter_by(email="excel@college.edu").one().roll_number == "1001"


def test_existing_student_without_a_batch_is_added_without_a_new_account(client, db, sent):
    admin = make_user(db, UserRole.admin)
    batch_id = _batch(db)
    loose = _student(db, None, urn="U1", email="loose@college.edu")
    body = _upload(client, admin, batch_id, "import", _csv([("U1", "Loose", "loose@college.edu")])).json()
    assert body["rows"][0]["status"] == "enrolled" and body["created"] == 0
    assert db.query(Enrollment).filter_by(student_id=loose.id, batch_id=batch_id).count() == 1
    assert sent == []


def test_bad_files_are_explained(client, db):
    admin = make_user(db, UserRole.admin)
    batch_id = _batch(db)
    assert _upload(client, admin, batch_id, "preview", b"hello", "notes.txt").status_code == 422
    assert _upload(client, admin, batch_id, "preview", _csv([("a", "b")], header=("Foo", "Bar", "Baz"))).status_code == 422
    assert _upload(client, admin, batch_id, "preview", _csv([])).status_code == 422
    assert _upload(client, admin, batch_id, "preview", b"not a real workbook", "x.xlsx").status_code == 422


def test_import_into_an_archived_batch_is_refused(client, db):
    admin = make_user(db, UserRole.admin)
    batch_id = _batch(db)
    client.post(f"/admin/batches/{batch_id}/archive", headers=admin.headers)
    assert _upload(client, admin, batch_id, "import", _csv([("U1", "A", "a@college.edu")])).status_code == 409


# ── Instructor console: Inactive students ─────────────────────────────────────

def test_deactivated_student_is_marked_inactive_in_the_instructor_console(client, db):
    admin = make_user(db, UserRole.admin)
    teacher = make_user(db, UserRole.instructor)
    batch_id = _batch(db)
    db.add(InstructorAssignment(instructor_id=teacher.id, batch_id=batch_id))
    student = _student(db, batch_id, name="Asha Patil", urn="CS001")
    exp = make_experiment(db)
    client.post(f"/admin/users/{student.id}/deactivate", headers=admin.headers)

    roster = client.get(f"/instructor/batches/{batch_id}", headers=teacher.headers).json()["students"]
    assert roster[0]["is_active"] is False  # still listed, but marked
    base = f"/instructor/batches/{batch_id}/experiments/{exp.id}"
    assert client.get(f"{base}/status", headers=teacher.headers).json()["rows"][0]["is_active"] is False
    assert client.get(f"{base}/assignment-status", headers=teacher.headers).json()["rows"][0]["is_active"] is False

    sheet = load_workbook(io.BytesIO(client.get(f"{base}/status/export", headers=teacher.headers).content)).active
    assert list(sheet.iter_rows(values_only=True))[6][1] == "Asha Patil (Inactive)"
